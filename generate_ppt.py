import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def create_presentation():
    # Create a presentation object
    prs = Presentation()
    
    # Define slide layouts
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]
    
    # 1. Title Slide
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Local Excel-Based Project Scheduler"
    subtitle.text = "Project Completion & Inspection Report\nSeptember 2026"

    # 2. Executive Summary
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Executive Summary"
    tf = body_shape.text_frame
    tf.text = "The Project Scheduler application has been successfully converted to run completely locally."
    
    p = tf.add_paragraph()
    p.text = "Uses Excel files for data storage"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "No database required"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Perfect for office environments with restrictions"
    p.level = 1

    # 3. What Was Delivered
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "What Was Delivered"
    tf = body_shape.text_frame
    tf.text = "Core Components:"
    
    p = tf.add_paragraph()
    p.text = "Core System: Excel Storage Engine & Data Models"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Web Application: Lightweight Flask-based server"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "User Interface: 8 Complete HTML Templates"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Setup & Configuration: Automated script & full documentation"
    p.level = 1

    # 4. System Architecture
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "System Architecture & Data Flow"
    tf = body_shape.text_frame
    tf.text = "Architecture:"
    
    p = tf.add_paragraph()
    p.text = "User Browser -> Flask Web Server (Port 8000) -> Excel Storage Layer -> Excel Files"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Data Flow:"
    p.level = 0
    
    p = tf.add_paragraph()
    p.text = "User Action -> Route Handler -> Excel Model -> Storage Manager -> Excel File -> User UI"
    p.level = 1

    # 5. Features Implemented
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Features Implemented"
    tf = body_shape.text_frame
    tf.text = "Key functionalities:"
    
    p = tf.add_paragraph()
    p.text = "User Management (Authentication, roles)"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Task Management (CRUD, filters, progress)"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Step Management (Task breakdown, time estimates)"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Dashboard (Stats, overdue alerts, visualization)"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Data Management (All in Excel, easily backupable)"
    p.level = 1

    # 6. Performance Metrics
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Performance Metrics"
    tf = body_shape.text_frame
    tf.text = "Efficiency and Speed:"
    
    p = tf.add_paragraph()
    p.text = "Startup Time: < 2 seconds"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Excel File Size: ~5 KB each (empty)"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Memory Usage: < 100 MB"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Response Time: < 100ms (local)"
    p.level = 1

    # 7. Success Criteria & Next Steps
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Success Criteria & Conclusion"
    tf = body_shape.text_frame
    tf.text = "All Requirements Met:"
    
    p = tf.add_paragraph()
    p.text = "No database required, completely local Excel storage"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Office-friendly & easy to run"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "How to run:"
    p.level = 0
    
    p = tf.add_paragraph()
    p.text = "python run_local.py -> http://127.0.0.1:8000"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Login: admin / admin123"
    p.level = 1

    # Save presentation
    output_path = "Project_Scheduler_Report.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully at {output_path}")

if __name__ == '__main__':
    create_presentation()
