import os
from pptx import Presentation
from pptx.util import Inches, Pt

def create_kaizen_presentation():
    prs = Presentation()
    
    title_slide_layout = prs.slide_layouts[0]
    content_slide_layout = prs.slide_layouts[1]
    title_only_layout = prs.slide_layouts[5]

    # --- Slide 1: Title & Theme ---
    slide1 = prs.slides.add_slide(title_slide_layout)
    title = slide1.shapes.title
    subtitle = slide1.placeholders[1]
    
    title.text = "Kaizen Report: Local Excel-Based Project Scheduler"
    subtitle.text = "Eliminating manual tracking and establishing a centralized platform for task and step monitoring."

    # --- Slide 2: Impact & Theme ---
    slide2 = prs.slides.add_slide(content_slide_layout)
    title2 = slide2.shapes.title
    title2.text = "Theme & Impact"
    
    tf2 = slide2.placeholders[1].text_frame
    tf2.text = "Theme: Implementation of a Local Excel-Based Project Scheduler to establish a centralized platform for project governance, task tracking, progress monitoring, and performance reporting."
    
    p = tf2.add_paragraph()
    p.text = "Impact:"
    p.level = 0
    
    impacts = [
        "1. Absence of centralized visibility over project activities and milestones.",
        "2. Difficulty in monitoring task timelines and step deliverables.",
        "3. Delays in tracking Turnaround Time (TAT) and overdue tasks.",
        "4. Increased effort scattering data across disparate files."
    ]
    for imp in impacts:
        p = tf2.add_paragraph()
        p.text = imp
        p.level = 1

    # --- Slide 3: Before Condition & View Point ---
    slide3 = prs.slides.add_slide(content_slide_layout)
    title3 = slide3.shapes.title
    title3.text = "Before Condition & Kaizen Thinking"
    
    tf3 = slide3.placeholders[1].text_frame
    tf3.text = "Before Condition:"
    
    before_conds = [
        "Task tracking was managed through disparate files and manual updates.",
        "No centralized dashboard or TAT calculation existed.",
        "Manual effort required for project status tracking and reporting."
    ]
    for bc in before_conds:
        p = tf3.add_paragraph()
        p.text = bc
        p.level = 1
        
    p = tf3.add_paragraph()
    p.text = "Process Flow (Before):"
    p.level = 0
    
    p = tf3.add_paragraph()
    p.text = "Manual Task Entry -> Unstructured Steps -> Manual Follow-ups -> No TAT Tracking -> Limited Visibility"
    p.level = 1
    
    p = tf3.add_paragraph()
    p.text = "View Point / Thinking of Kaizen:"
    p.level = 0
    
    kaizen_thoughts = [
        "To create a localized, database-free, Excel-backend system.",
        "To enforce structured task breakdown and dependencies.",
        "To enable automatic TAT tracking and overdue notifications.",
        "To support effective project review through dashboard-based reporting."
    ]
    for kt in kaizen_thoughts:
        p = tf3.add_paragraph()
        p.text = kt
        p.level = 1

    # --- Slide 4: After Condition ---
    slide4 = prs.slides.add_slide(content_slide_layout)
    title4 = slide4.shapes.title
    title4.text = "After Condition"
    
    tf4 = slide4.placeholders[1].text_frame
    tf4.text = "After Condition:"
    
    after_conds = [
        "1. A lightweight Flask web application created for centralized task management.",
        "2. Clear visibility of overall and stage-wise progress via dashboard.",
        "3. Automatic email auto-triggers and overdue alerts configured.",
        "4. Seamless data persistence entirely in local Excel sheets."
    ]
    for ac in after_conds:
        p = tf4.add_paragraph()
        p.text = ac
        p.level = 1
        
    p = tf4.add_paragraph()
    p.text = "Process Flow (After):"
    p.level = 0
    
    p = tf4.add_paragraph()
    p.text = "Task Creation -> Step Breakdown & Dependencies -> Automated TAT Tracking -> Auto-Sync to Excel -> Real-Time Dashboard Visibility"
    p.level = 1

    # --- Slide 5: Benefits & Cost Savings Table ---
    slide5 = prs.slides.add_slide(title_only_layout)
    title5 = slide5.shapes.title
    title5.text = "Benefits & Time Savings"
    
    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(4.5)
    height = Inches(2.0)
    
    # Text box for Benefits
    txBox = slide5.shapes.add_textbox(left, top, width, height)
    tf5 = txBox.text_frame
    tf5.text = "Benefits:"
    
    benefits = [
        "Visibility: Single-view dashboard for all tasks and sub-steps.",
        "Decision Making: Automatic overdue alerts and progress visualization.",
        "Productivity: Zero database setup required, easy local deployment, and reduced manual tracking effort."
    ]
    for b in benefits:
        p = tf5.add_paragraph()
        p.text = b
        p.level = 1
        
    # Table for metrics
    left_table = Inches(5.0)
    top_table = Inches(1.5)
    width_table = Inches(4.5)
    height_table = Inches(1.5)
    
    table_shape = slide5.shapes.add_table(5, 3, left_table, top_table, width_table, height_table)
    table = table_shape.table
    
    # Headers
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(0, 2).text = "Calculation"
    
    # Row 1
    table.cell(1, 0).text = "Time saved per month"
    table.cell(1, 1).text = "132 mins"
    table.cell(1, 2).text = "132/60 = 2.2 hrs"
    
    # Row 2
    table.cell(2, 0).text = "Cost per hour"
    table.cell(2, 1).text = "₹300/hour"
    table.cell(2, 2).text = "₹300/hour"
    
    # Row 3
    table.cell(3, 0).text = "Monthly cost saving"
    table.cell(3, 1).text = "₹660/month"
    table.cell(3, 2).text = "2.2 x 300 = ₹660"
    
    # Row 4
    table.cell(4, 0).text = "Annual cost saving"
    table.cell(4, 1).text = "₹7,920/year"
    table.cell(4, 2).text = "660 x 12 = ₹7,920/year"

    # --- Slide 6: Standardization ---
    slide6 = prs.slides.add_slide(content_slide_layout)
    title6 = slide6.shapes.title
    title6.text = "Standardization"
    
    tf6 = slide6.placeholders[1].text_frame
    tf6.text = "Standardization Rules:"
    
    stds = [
        "1. All task statuses and steps to be updated exclusively via the web UI.",
        "2. Standard process defined for TAT calculation across dependencies.",
        "3. Web Dashboard to be used as the single source for review and monitoring.",
        "4. Excel backend files are locked from manual edits to prevent data corruption."
    ]
    for st in stds:
        p = tf6.add_paragraph()
        p.text = st
        p.level = 1

    # Save presentation
    output_path = "Kaizen_Project_Scheduler_Adapted.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully at {output_path}")

if __name__ == '__main__':
    create_kaizen_presentation()
